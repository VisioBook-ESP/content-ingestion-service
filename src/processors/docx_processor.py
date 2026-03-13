from src.processors.base import BaseProcessor

try:
    from docx import Document

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class DocxProcessor(BaseProcessor):
    def can_handle(self, file_path: str) -> bool:
        return file_path.lower().endswith(".docx")

    def extract_text(self, file_path: str) -> str:
        if not DOCX_AVAILABLE:
            raise ImportError(
                "python-docx is required for DOCX processing. "
                "Install it with: pip install python-docx"
            )

        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)

    def extract_metadata(self, file_path: str) -> dict:
        if not DOCX_AVAILABLE:
            return {"title": None, "author": None, "pages": None}

        doc = Document(file_path)
        core_props = doc.core_properties

        return {
            "title": core_props.title or None,
            "author": core_props.author or None,
            "pages": None,
            "created": core_props.created.isoformat() if core_props.created else None,
            "modified": core_props.modified.isoformat() if core_props.modified else None,
            "subject": core_props.subject or None,
        }
